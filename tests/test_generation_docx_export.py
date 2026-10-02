from io import BytesIO

from docx import Document
from docx.document import Document as DocxDocument

from sbobina.docx_export import (
    render_exam_docx,
    render_solutions_docx,
    render_summary_docx,
)
from sbobina.generation_models import (
    GenerationCitation,
    GenerationFormat,
    GenerationQuestion,
    GenerationRecord,
    GenerationStatus,
    SummarySection,
    SummarySentence,
)

OPTIONS = ("Opzione A", "Opzione B", "Opzione C", "Opzione D")


def _open(data: bytes) -> DocxDocument:
    return Document(BytesIO(data))


def _mc_question(
    question: str = "Quale regola si applica?",
    correct_index: int = 2,
    citations: tuple[GenerationCitation, ...] = (),
) -> GenerationQuestion:
    return GenerationQuestion(
        question=question,
        options=OPTIONS,
        correct_index=correct_index,
        solution="La regola corretta è la C per via dell'articolo 1321.",
        citations=citations,
    )


def _open_question(
    question: str = "Spiega il principio.",
    citations: tuple[GenerationCitation, ...] = (),
) -> GenerationQuestion:
    return GenerationQuestion(
        question=question,
        options=(),
        correct_index=None,
        solution="Il principio richiede X, Y, Z.",
        citations=citations,
    )


def _record(
    format_: GenerationFormat = GenerationFormat.MULTIPLE_CHOICE,
    topic: str = "Diritto privato",
    questions: tuple[GenerationQuestion, ...] = (),
    sections: tuple[SummarySection, ...] = (),
) -> GenerationRecord:
    return GenerationRecord(
        id="gen-1",
        format=format_,
        status=GenerationStatus.DONE,
        requested_count=len(questions) or 1,
        topic=topic,
        model="qwen3.5:9b",
        prompt_version="v1",
        generated_at="2026-10-03T10:00:00Z",
        sources=(),
        discarded=(),
        questions=questions,
        sections=sections,
        error=None,
    )


# --- compito.docx ---


def _doc_citation(quote: str, page: int = 1) -> GenerationCitation:
    return GenerationCitation(
        passage_id=f"d1:p{page}:c0",
        quote=quote,
        doc_id="d1",
        page=page,
        job_id=None,
        timestamp=None,
    )


def test_render_exam_docx_numbers_questions_with_lettered_options() -> None:
    questions = (
        _mc_question(question="Prima domanda?"),
        _mc_question(question="Seconda domanda?"),
    )
    document = _open(render_exam_docx(_record(questions=questions)))

    texts = [p.text for p in document.paragraphs]
    assert any(t.startswith("1. Prima domanda?") for t in texts)
    assert any(t.startswith("2. Seconda domanda?") for t in texts)
    assert "a) Opzione A" in texts
    assert "d) Opzione D" in texts


def test_render_exam_docx_never_leaks_solution_or_correct_option() -> None:
    question = _mc_question(
        question="Domanda riservata?",
        citations=(_doc_citation(quote="testo segreto"),),
    )
    document = _open(render_exam_docx(_record(questions=(question,))))

    full_text = "\n".join(p.text for p in document.paragraphs)
    assert question.solution not in full_text
    assert "testo segreto" not in full_text
    assert "Risposta corretta" not in full_text
    assert "corretta" not in full_text.lower()


def test_render_exam_docx_open_question_has_no_lettered_options() -> None:
    document = _open(
        render_exam_docx(
            _record(format_=GenerationFormat.OPEN, questions=(_open_question(),))
        )
    )

    texts = [p.text for p in document.paragraphs]
    assert not any(t.startswith("a) ") for t in texts)


def test_render_exam_docx_uses_topic_as_title() -> None:
    document = _open(
        render_exam_docx(_record(topic="Diritto privato", questions=(_mc_question(),)))
    )

    assert document.paragraphs[0].text == "Diritto privato"
    assert document.core_properties.title == "Diritto privato"


def test_render_exam_docx_falls_back_to_generic_title_when_topic_is_blank() -> None:
    document = _open(render_exam_docx(_record(topic="", questions=(_mc_question(),))))

    assert document.paragraphs[0].text == "Compito"


# --- soluzioni.docx ---


def test_render_solutions_docx_includes_solution_text_and_correct_letter() -> None:
    question = _mc_question(question="Domanda?", correct_index=1)
    document = _open(render_solutions_docx(_record(questions=(question,))))

    full_text = "\n".join(p.text for p in document.paragraphs)
    assert question.solution in full_text
    assert "Risposta corretta: b" in full_text


def test_render_solutions_docx_formats_citation_with_quote_and_reference() -> None:
    question = _mc_question(citations=(_doc_citation(quote="testo citato", page=3),))
    document = _open(
        render_solutions_docx(
            _record(questions=(question,)), doc_filenames={"d1": "Manuale.pdf"}
        )
    )

    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "«testo citato» (Manuale.pdf, pagina 3)" in full_text


def test_render_solutions_docx_open_question_has_no_correct_letter_line() -> None:
    document = _open(
        render_solutions_docx(
            _record(format_=GenerationFormat.OPEN, questions=(_open_question(),))
        )
    )

    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "Risposta corretta" not in full_text
    assert "Il principio richiede X, Y, Z." in full_text


def test_render_solutions_docx_is_a_separate_file_from_the_exam() -> None:
    question = _mc_question()
    record = _record(questions=(question,))

    exam_text = "\n".join(p.text for p in _open(render_exam_docx(record)).paragraphs)
    solutions_text = "\n".join(
        p.text for p in _open(render_solutions_docx(record)).paragraphs
    )

    assert question.solution not in exam_text
    assert question.solution in solutions_text


# --- riassunto.docx ---


def test_render_summary_docx_sections_as_headings_sentences_as_paragraphs() -> None:
    section = SummarySection(
        title="Le obbligazioni",
        sentences=(
            SummarySentence(text="Prima frase.", citations=()),
            SummarySentence(text="Seconda frase.", citations=()),
        ),
    )
    document = _open(
        render_summary_docx(
            _record(
                format_=GenerationFormat.SUMMARY, topic="Corso", sections=(section,)
            )
        )
    )

    texts = [p.text for p in document.paragraphs]
    assert "Le obbligazioni" in texts
    assert "Prima frase." in texts
    assert "Seconda frase." in texts
    heading_paragraph = next(
        p for p in document.paragraphs if p.text == "Le obbligazioni"
    )
    assert heading_paragraph.style is not None
    assert heading_paragraph.style.name.startswith("Heading")


def test_render_summary_docx_appends_citation_reference_to_sentence() -> None:
    sentence = SummarySentence(
        text="Il contratto vincola le parti.",
        citations=(_doc_citation(quote="irrelevant", page=2),),
    )
    section = SummarySection(title="Contratti", sentences=(sentence,))
    document = _open(
        render_summary_docx(
            _record(format_=GenerationFormat.SUMMARY, sections=(section,)),
            doc_filenames={"d1": "Manuale.pdf"},
        )
    )

    texts = [p.text for p in document.paragraphs]
    assert "Il contratto vincola le parti. (Manuale.pdf, pagina 2)" in texts


def test_render_summary_docx_falls_back_to_generic_title_when_topic_is_blank() -> None:
    section = SummarySection(
        title="X", sentences=(SummarySentence(text="Y.", citations=()),)
    )
    document = _open(
        render_summary_docx(
            _record(format_=GenerationFormat.SUMMARY, topic="", sections=(section,))
        )
    )

    assert document.paragraphs[0].text == "Riassunto"
