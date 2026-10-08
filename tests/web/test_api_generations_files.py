from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from generation_api_fixtures import (
    COURSES_URL,
    _make_record,
    _mc_question,
    _register_course,
    _write_document,
    client,
)

from sbobina.generation_models import (
    GenerationCitation,
    GenerationFormat,
    GenerationQuestion,
    GenerationRecord,
    GenerationSourceUsed,
    GenerationStatus,
    SummarySection,
    SummarySentence,
)

__all__ = ["client"]


def test_download_compito_md_never_contains_solution_text(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        status=GenerationStatus.DONE,
        questions=(_mc_question(correct="la soluzione segreta"),),
    )

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/compito.md"
    )

    assert response.status_code == 200
    assert "soluzione segreta" not in response.text
    assert record.questions[0].question in response.text
    solutions = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/soluzioni.md"
    )
    assert solutions.status_code == 200
    assert "soluzione segreta" in solutions.text
    assert response.headers["x-content-type-options"] == "nosniff"
    assert 'filename="compito.md"' in response.headers["content-disposition"]


def test_download_soluzioni_md_contains_solution(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        status=GenerationStatus.DONE,
        questions=(_mc_question(correct="la soluzione segreta"),),
    )

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/soluzioni.md"
    )

    assert response.status_code == 200
    assert "la soluzione segreta" in response.text


def test_download_generation_file_wrong_format_is_not_found(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        status=GenerationStatus.DONE,
        questions=(_mc_question(correct="x"),),
    )

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/riassunto.md"
    )

    assert response.status_code == 404


def test_download_generation_file_unknown_name_is_not_found(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path, course_id=course_id, status=GenerationStatus.DONE
    )

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/segreto.txt"
    )

    assert response.status_code == 404


def test_download_generation_file_not_done_is_not_found(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(tmp_path=tmp_path, course_id=course_id)

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/compito.md"
    )

    assert response.status_code == 404


def test_download_riassunto_docx(client: TestClient, tmp_path: Path) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    section = SummarySection(
        title="Introduzione",
        sentences=(
            SummarySentence(text="La forza e' massa per accelerazione.", citations=()),
        ),
    )
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        format_=GenerationFormat.SUMMARY,
        status=GenerationStatus.DONE,
        sections=(section,),
    )

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/riassunto.docx"
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    paragraphs = [
        paragraph.text for paragraph in Document(BytesIO(response.content)).paragraphs
    ]
    assert section.title in paragraphs
    assert section.sentences[0].text in paragraphs


def test_download_compito_docx_never_contains_solution_text(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        status=GenerationStatus.DONE,
        questions=(_mc_question(correct="la soluzione segreta"),),
    )
    base = f"{COURSES_URL}/fisica/generations/{record.id}/files"

    exam = client.get(f"{base}/compito.docx")
    solutions = client.get(f"{base}/soluzioni.docx")

    def text(data: bytes) -> str:
        return "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)

    assert exam.status_code == 200
    assert "soluzione segreta" not in text(exam.content)
    assert "soluzione segreta" in text(solutions.content)


def _document_citation(doc_id: str, page: int) -> GenerationCitation:
    return GenerationCitation(
        passage_id=f"{doc_id}:p{page}:c0",
        quote="la forza e massa per accelerazione",
        doc_id=doc_id,
        page=page,
        job_id=None,
        timestamp=None,
    )


def _document_source(doc_id: str) -> GenerationSourceUsed:
    return GenerationSourceUsed(
        doc_id=doc_id, sha256="a" * 64, job_id=None, revision=None
    )


def _open_record_citing(
    tmp_path: Path, course_id: str, doc_id: str
) -> GenerationRecord:
    question = GenerationQuestion(
        question="Che cos'e' la forza?",
        options=(),
        correct_index=None,
        solution="massa per accelerazione",
        citations=(_document_citation(doc_id=doc_id, page=12),),
    )
    return _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        format_=GenerationFormat.OPEN,
        status=GenerationStatus.DONE,
        questions=(question,),
        sources=(
            _document_source(doc_id=doc_id),
            GenerationSourceUsed(doc_id=None, sha256=None, job_id="j1", revision="r1"),
        ),
    )


def test_download_soluzioni_md_names_cited_document_by_current_filename(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    _write_document(
        tmp_path=tmp_path, course_id=course_id, doc_id="doc-1", filename="Manuale.pdf"
    )
    record = _open_record_citing(tmp_path=tmp_path, course_id=course_id, doc_id="doc-1")

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/soluzioni.md"
    )

    assert "(Manuale.pdf, pagina 12)" in response.text


def test_download_soluzioni_md_removed_document_falls_back_to_its_id(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _open_record_citing(tmp_path=tmp_path, course_id=course_id, doc_id="doc-9")

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/soluzioni.md"
    )

    assert response.status_code == 200
    assert "(doc-9, pagina 12)" in response.text


def test_download_soluzioni_md_open_question_has_no_correct_option_line(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    open_record = _open_record_citing(
        tmp_path=tmp_path, course_id=course_id, doc_id="doc-1"
    )
    mc_record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        status=GenerationStatus.DONE,
        questions=(_mc_question(correct="x"),),
    )

    def soluzioni(record_id: str) -> str:
        url = f"{COURSES_URL}/fisica/generations/{record_id}/files/soluzioni.md"
        text: str = client.get(url).text
        return text

    assert "Risposta corretta" not in soluzioni(open_record.id)
    assert "Risposta corretta: a) a" in soluzioni(mc_record.id)


def test_download_riassunto_md_lists_sections_sentences_and_sources(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    _write_document(
        tmp_path=tmp_path, course_id=course_id, doc_id="doc-1", filename="Manuale.pdf"
    )
    sentence = SummarySentence(
        text="La forza e' massa per accelerazione.",
        citations=(_document_citation(doc_id="doc-1", page=3),),
    )
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        format_=GenerationFormat.SUMMARY,
        status=GenerationStatus.DONE,
        sections=(SummarySection(title="Dinamica", sentences=(sentence,)),),
        sources=(_document_source(doc_id="doc-1"),),
    )

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/riassunto.md"
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/markdown")
    expected = ("Dinamica", "massa per accelerazione", "Manuale.pdf, pagina 3")
    assert [fragment for fragment in expected if fragment not in response.text] == []


def test_get_generation_summary_resolves_sentence_citations(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    _write_document(
        tmp_path=tmp_path, course_id=course_id, doc_id="doc-1", filename="Manuale.pdf"
    )
    sentence = SummarySentence(
        text="La forza e' massa per accelerazione.",
        citations=(_document_citation(doc_id="doc-1", page=3),),
    )
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        format_=GenerationFormat.SUMMARY,
        status=GenerationStatus.DONE,
        sections=(SummarySection(title="Dinamica", sentences=(sentence,)),),
        sources=(_document_source(doc_id="doc-1"),),
    )

    response = client.get(f"{COURSES_URL}/fisica/generations/{record.id}")
    [section] = response.json()["data"]["sections"]
    [payload] = section["sentences"]

    assert section["title"] == "Dinamica"
    assert payload["text"] == "La forza e' massa per accelerazione."
    assert (payload["citations"][0]["source"], payload["citations"][0]["href"]) == (
        "Manuale.pdf",
        "/corsi/fisica/documenti/doc-1?p=3",
    )


def test_get_generation_unregistered_course_is_not_found(client: TestClient) -> None:
    response = client.get(f"{COURSES_URL}/sconosciuto/generations/qualunque")

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "NOT_FOUND"


def test_list_generations_registered_course_without_generations_is_empty(
    client: TestClient, tmp_path: Path
) -> None:
    _register_course(tmp_path=tmp_path)

    response = client.get(f"{COURSES_URL}/fisica/generations")

    assert response.json()["data"] == []
