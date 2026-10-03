from collections.abc import Iterator
from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from sbobina.course_registry import get_or_create
from sbobina.document_models import CourseDocument, DocumentKind, DocumentStatus
from sbobina.generation_models import (
    GenerationCitation,
    GenerationFormat,
    GenerationQuestion,
    GenerationRecord,
    GenerationRequest,
    GenerationSourceUsed,
    GenerationStatus,
    SummarySection,
    SummarySentence,
)
from sbobina.models import Segment, Transcript, Word, save_transcript
from sbobina.settings import Settings
from sbobina.web.api_files import TRANSCRIPT_FILES
from sbobina.web.app import create_app
from sbobina.web.document_store import document_dir, write_document
from sbobina.web.generation_store import (
    create_generation,
    generation_path,
    save_generation,
)
from sbobina.web.job_store import JobStore

BASE_URL = "http://127.0.0.1:8765"
COURSES_URL = "/api/v1/courses"


@pytest.fixture
def client(tmp_path: Path) -> Iterator[TestClient]:
    app = create_app(settings=Settings(), data_dir=tmp_path)
    transport = TestClient(app=app, base_url=BASE_URL, headers={"Origin": BASE_URL})
    yield transport
    transport.close()


def _store(tmp_path: Path) -> JobStore:
    return JobStore(data_dir=tmp_path)


def _register_course(tmp_path: Path, key: str = "fisica") -> str:
    course = get_or_create(
        courses_dir=_store(tmp_path).courses_dir, key=key, label=key.title()
    )
    return course.id


def _write_document(tmp_path: Path, course_id: str, doc_id: str, filename: str) -> None:
    courses_dir = _store(tmp_path).courses_dir
    doc_dir = document_dir(courses_dir=courses_dir, course_id=course_id, doc_id=doc_id)
    doc_dir.mkdir(parents=True)
    write_document(
        courses_dir=courses_dir,
        document=CourseDocument(
            id=doc_id,
            course_id=course_id,
            filename=filename,
            kind=DocumentKind.PDF,
            size=10,
            sha256="a" * 64,
            status=DocumentStatus.READY,
            error=None,
            pages=1,
            created_at=datetime.now(UTC),
        ),
    )


def _write_lecture(tmp_path: Path, job_id: str) -> None:
    store = _store(tmp_path)
    transcript = Transcript(
        source="lezione.m4a",
        model="large-v3",
        language="it",
        duration=10.0,
        segments=(
            Segment(
                start=0.0,
                end=6.0,
                words=(
                    Word(start=0.0, end=1.0, text="il ", probability=0.99),
                    Word(start=1.0, end=2.0, text="gatto ", probability=0.99),
                    Word(start=2.0, end=3.0, text="nero ", probability=0.99),
                    Word(start=3.0, end=4.0, text="dorme ", probability=0.99),
                    Word(start=4.0, end=5.0, text="sul ", probability=0.99),
                    Word(start=5.0, end=6.0, text="tappeto", probability=0.99),
                ),
            ),
        ),
    )
    directory = store.jobs_dir / job_id
    directory.mkdir(parents=True, exist_ok=True)
    save_transcript(
        transcript=transcript, path=directory / TRANSCRIPT_FILES["original"]
    )


def _make_record(
    tmp_path: Path,
    course_id: str,
    *,
    format_: GenerationFormat = GenerationFormat.MULTIPLE_CHOICE,
    status: GenerationStatus = GenerationStatus.QUEUED,
    questions: tuple[GenerationQuestion, ...] = (),
    sections: tuple[SummarySection, ...] = (),
    sources: tuple[GenerationSourceUsed, ...] = (),
) -> GenerationRecord:
    courses_dir = _store(tmp_path).courses_dir
    base = create_generation(
        courses_dir=courses_dir,
        course_id=course_id,
        request=GenerationRequest(format=format_, count=1),
    )
    is_done = status == GenerationStatus.DONE
    record = replace(
        base,
        status=status,
        questions=questions,
        sections=sections,
        sources=sources,
        model="qwen-test" if is_done else "",
        prompt_version="v1" if is_done else "",
        generated_at="2026-01-01T00:00:00" if is_done else "",
    )
    save_generation(courses_dir=courses_dir, course_id=course_id, record=record)
    return record


def test_create_generation_accepted_and_queued(
    client: TestClient, tmp_path: Path
) -> None:
    _register_course(tmp_path=tmp_path)
    response = client.post(
        f"{COURSES_URL}/fisica/generations",
        json={"format": "multiple_choice", "count": 5, "topic": "cinematica"},
    )
    assert response.status_code == 202
    data = response.json()["data"]
    assert data["status"] == "queued"
    assert data["requested_count"] == 5


def test_create_generation_validation_error(client: TestClient, tmp_path: Path) -> None:
    _register_course(tmp_path=tmp_path)
    response = client.post(
        f"{COURSES_URL}/fisica/generations",
        json={"format": "multiple_choice", "count": 0},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "VALIDATION_ERROR"


def test_create_generation_empty_course_key_is_conflict(client: TestClient) -> None:
    response = client.post(
        f"{COURSES_URL}//generations", json={"format": "open", "count": 1}
    )
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "COURSE_REQUIRED"


def test_create_generation_unregistered_course_is_not_found(client: TestClient) -> None:
    response = client.post(
        f"{COURSES_URL}/sconosciuto/generations", json={"format": "open", "count": 1}
    )
    assert response.status_code == 404


def test_create_generation_foreign_origin_rejected(
    client: TestClient, tmp_path: Path
) -> None:
    _register_course(tmp_path=tmp_path)
    response = client.post(
        f"{COURSES_URL}/fisica/generations",
        json={"format": "open", "count": 1},
        headers={"Origin": "http://evil.example"},
    )
    assert response.status_code == 403
    assert response.json()["error"]["code"] == "FORBIDDEN"


def test_list_generations_paginated_newest_first(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    first = _make_record(tmp_path=tmp_path, course_id=course_id)
    second = _make_record(tmp_path=tmp_path, course_id=course_id)
    generation_path(
        courses_dir=_store(tmp_path).courses_dir, course_id=course_id, gen_id=second.id
    ).touch()

    response = client.get(f"{COURSES_URL}/fisica/generations", params={"per_page": 1})
    body = response.json()

    assert response.status_code == 200
    assert body["meta"] == {"page": 1, "per_page": 1, "total": 2, "total_pages": 2}
    assert body["data"][0]["id"] == second.id
    second_page = client.get(
        f"{COURSES_URL}/fisica/generations", params={"page": 2, "per_page": 1}
    ).json()
    assert second_page["data"][0]["id"] == first.id


def test_list_generations_unregistered_course_is_empty(client: TestClient) -> None:
    response = client.get(f"{COURSES_URL}/sconosciuto/generations")
    assert response.status_code == 200
    assert response.json() == {
        "data": [],
        "meta": {"page": 1, "per_page": 20, "total": 0, "total_pages": 0},
    }


def test_get_generation_not_found(client: TestClient, tmp_path: Path) -> None:
    _register_course(tmp_path=tmp_path)
    response = client.get(f"{COURSES_URL}/fisica/generations/missing")
    assert response.status_code == 404


def test_get_generation_resolves_document_citation(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    _write_document(
        tmp_path=tmp_path, course_id=course_id, doc_id="doc-1", filename="Manuale.pdf"
    )
    question = GenerationQuestion(
        question="Che cos'e' la forza?",
        options=(),
        correct_index=None,
        solution="massa per accelerazione",
        citations=(
            GenerationCitation(
                passage_id="D1",
                quote="la forza e massa per accelerazione",
                doc_id="doc-1",
                page=12,
                job_id=None,
                timestamp=None,
            ),
        ),
    )
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        format_=GenerationFormat.OPEN,
        status=GenerationStatus.DONE,
        questions=(question,),
        sources=(
            GenerationSourceUsed(
                doc_id="doc-1", sha256="a" * 64, job_id=None, revision=None
            ),
        ),
    )

    response = client.get(f"{COURSES_URL}/fisica/generations/{record.id}")
    citation = response.json()["data"]["questions"][0]["citations"][0]

    assert response.status_code == 200
    assert citation["source"] == "Manuale.pdf"
    assert citation["href"] == "/corsi/fisica/documenti/doc-1?p=12"


def test_get_generation_document_citation_removed_source(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    question = GenerationQuestion(
        question="Domanda",
        options=(),
        correct_index=None,
        solution="risposta",
        citations=(
            GenerationCitation(
                passage_id="D1",
                quote="testo citato ma non piu disponibile",
                doc_id="doc-mancante",
                page=3,
                job_id=None,
                timestamp=None,
            ),
        ),
    )
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        format_=GenerationFormat.OPEN,
        status=GenerationStatus.DONE,
        questions=(question,),
    )

    response = client.get(f"{COURSES_URL}/fisica/generations/{record.id}")
    citation = response.json()["data"]["questions"][0]["citations"][0]

    assert citation["source"] == "fonte rimossa"
    assert citation["href"] is None


def test_get_generation_resolves_lecture_citation_exact_timestamp(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    _write_lecture(tmp_path=tmp_path, job_id="lezione-1")
    question = GenerationQuestion(
        question="Di che colore e' il gatto?",
        options=(),
        correct_index=None,
        solution="nero",
        citations=(
            GenerationCitation(
                passage_id="L1",
                quote="dorme sul tappeto",
                doc_id=None,
                page=None,
                job_id="lezione-1",
                timestamp=0.0,
            ),
        ),
    )
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        format_=GenerationFormat.OPEN,
        status=GenerationStatus.DONE,
        questions=(question,),
    )

    response = client.get(f"{COURSES_URL}/fisica/generations/{record.id}")
    citation = response.json()["data"]["questions"][0]["citations"][0]

    # The anchor is the segment start (0.0); the cited words start later, at
    # the "dorme" word (3.0s): the resolved timestamp must be the exact one.
    assert citation["timestamp"] == 3.0
    assert citation["href"] == "/lettore/lezione-1?t=3.0&variant=original"


def test_get_generation_lecture_citation_falls_back_to_anchor(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    _write_lecture(tmp_path=tmp_path, job_id="lezione-1")
    question = GenerationQuestion(
        question="Domanda",
        options=(),
        correct_index=None,
        solution="risposta",
        citations=(
            GenerationCitation(
                passage_id="L1",
                quote="frase che non esiste piu nella trascrizione",
                doc_id=None,
                page=None,
                job_id="lezione-1",
                timestamp=0.0,
            ),
        ),
    )
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        format_=GenerationFormat.OPEN,
        status=GenerationStatus.DONE,
        questions=(question,),
    )

    response = client.get(f"{COURSES_URL}/fisica/generations/{record.id}")
    citation = response.json()["data"]["questions"][0]["citations"][0]

    assert citation["timestamp"] == 0.0


def _mc_question(correct: str) -> GenerationQuestion:
    return GenerationQuestion(
        question="Quale e' corretta?",
        options=("a", "b", "c", "d"),
        correct_index=0,
        solution=correct,
        citations=(),
    )


def test_download_compito_md_never_contains_solution_text(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        status=GenerationStatus.DONE,
        questions=(_mc_question(correct="la soluzione segreta"),),
    )

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/compito.md"
    )

    assert response.status_code == 200
    assert "soluzione segreta" not in response.text
    assert response.headers["x-content-type-options"] == "nosniff"
    assert 'filename="compito.md"' in response.headers["content-disposition"]


def test_download_soluzioni_md_contains_solution(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        status=GenerationStatus.DONE,
        questions=(_mc_question(correct="la soluzione segreta"),),
    )

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/soluzioni.md"
    )

    assert response.status_code == 200
    assert "la soluzione segreta" in response.text


def test_download_generation_file_wrong_format_is_not_found(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        status=GenerationStatus.DONE,
        questions=(_mc_question(correct="x"),),
    )

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/riassunto.md"
    )

    assert response.status_code == 404


def test_download_generation_file_unknown_name_is_not_found(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path, course_id=course_id, status=GenerationStatus.DONE
    )

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/segreto.txt"
    )

    assert response.status_code == 404


def test_download_generation_file_not_done_is_not_found(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(tmp_path=tmp_path, course_id=course_id)

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/compito.md"
    )

    assert response.status_code == 404


def test_download_riassunto_docx(client: TestClient, tmp_path: Path) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    section = SummarySection(
        title="Introduzione",
        sentences=(
            SummarySentence(text="La forza e' massa per accelerazione.", citations=()),
        ),
    )
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        format_=GenerationFormat.SUMMARY,
        status=GenerationStatus.DONE,
        sections=(section,),
    )

    response = client.get(
        f"{COURSES_URL}/fisica/generations/{record.id}/files/riassunto.docx"
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def test_delete_generation_busy_while_queued(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(tmp_path=tmp_path, course_id=course_id)

    response = client.delete(f"{COURSES_URL}/fisica/generations/{record.id}")

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "GENERATION_BUSY"


def test_delete_generation_removes_it(client: TestClient, tmp_path: Path) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path, course_id=course_id, status=GenerationStatus.DONE
    )

    response = client.delete(f"{COURSES_URL}/fisica/generations/{record.id}")

    assert response.status_code == 204
    assert (
        client.get(f"{COURSES_URL}/fisica/generations/{record.id}").status_code == 404
    )


def test_delete_generation_foreign_origin_rejected(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path, course_id=course_id, status=GenerationStatus.DONE
    )

    response = client.delete(
        f"{COURSES_URL}/fisica/generations/{record.id}",
        headers={"Origin": "http://evil.example"},
    )

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "FORBIDDEN"


def test_cancel_generation_queued(client: TestClient, tmp_path: Path) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(tmp_path=tmp_path, course_id=course_id)

    response = client.post(f"{COURSES_URL}/fisica/generations/{record.id}/cancel")

    assert response.status_code == 200
    assert response.json()["data"]["status"] == "interrupted"


def test_cancel_generation_not_cancellable_is_conflict(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path, course_id=course_id, status=GenerationStatus.DONE
    )

    response = client.post(f"{COURSES_URL}/fisica/generations/{record.id}/cancel")

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "JOB_NOT_CANCELLABLE"


def test_get_generation_lecture_citation_removed_source(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    question = GenerationQuestion(
        question="Domanda",
        options=(),
        correct_index=None,
        solution="risposta",
        citations=(
            GenerationCitation(
                passage_id="L1",
                quote="testo di una lezione poi cancellata",
                doc_id=None,
                page=None,
                job_id="lezione-cancellata",
                timestamp=12.0,
            ),
        ),
    )
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        format_=GenerationFormat.OPEN,
        status=GenerationStatus.DONE,
        questions=(question,),
    )

    response = client.get(f"{COURSES_URL}/fisica/generations/{record.id}")
    citation = response.json()["data"]["questions"][0]["citations"][0]

    assert citation["source"] == "fonte rimossa"
    assert citation["href"] is None


def test_download_compito_docx_never_contains_solution_text(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        status=GenerationStatus.DONE,
        questions=(_mc_question(correct="la soluzione segreta"),),
    )
    base = f"{COURSES_URL}/fisica/generations/{record.id}/files"

    exam = client.get(f"{base}/compito.docx")
    solutions = client.get(f"{base}/soluzioni.docx")

    def text(data: bytes) -> str:
        return "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)

    assert exam.status_code == 200
    assert "soluzione segreta" not in text(exam.content)
    assert "soluzione segreta" in text(solutions.content)


def test_delete_generation_twice_is_not_a_server_error(
    client: TestClient, tmp_path: Path
) -> None:
    course_id = _register_course(tmp_path=tmp_path)
    record = _make_record(
        tmp_path=tmp_path,
        course_id=course_id,
        status=GenerationStatus.DONE,
        questions=(_mc_question(correct="x"),),
    )
    url = f"{COURSES_URL}/fisica/generations/{record.id}"

    first = client.delete(url)
    second = client.delete(url)

    assert first.status_code == 204
    assert second.status_code == 404
