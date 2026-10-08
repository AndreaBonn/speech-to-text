from io import BytesIO

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm

from sbobina.docx_export import render_book_docx

A4_WIDTH_CM = 21.0
A4_HEIGHT_CM = 29.7


def _open(data: bytes) -> DocxDocument:
    return Document(BytesIO(data))


def test_render_book_docx_title_then_one_paragraph_per_block() -> None:
    document = _open(
        render_book_docx(title="Diritto privato", paragraphs=["Primo.", "Secondo."])
    )

    texts = [p.text for p in document.paragraphs]
    assert texts == ["Diritto privato", "Primo.", "Secondo."]
    title_style = document.paragraphs[0].style
    assert title_style is not None
    assert title_style.name == "Title"
    assert document.core_properties.title == "Diritto privato"


def test_render_book_docx_body_is_justified_with_first_line_indent() -> None:
    document = _open(render_book_docx(title="T", paragraphs=["Uno.", "Due."]))

    body = document.styles["Normal"]
    assert body.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    indent = body.paragraph_format.first_line_indent
    assert indent is not None
    # Word stores the indent in twips, so 0.6 cm comes back as 0.5997 cm.
    assert indent.cm == pytest.approx(0.6, abs=0.01)


def test_render_book_docx_page_is_a4() -> None:
    section = _open(render_book_docx(title="T", paragraphs=["x"])).sections[0]

    assert section.page_width is not None
    assert section.page_height is not None
    assert round(section.page_width.cm, 1) == round(Cm(A4_WIDTH_CM).cm, 1)
    assert round(section.page_height.cm, 1) == round(Cm(A4_HEIGHT_CM).cm, 1)


def test_render_book_docx_footer_has_page_number_field() -> None:
    section = _open(render_book_docx(title="T", paragraphs=["x"])).sections[0]

    fields = section.footer.paragraphs[0]._p.findall(path=qn(tag="w:fldSimple"))
    assert [field.get(key=qn(tag="w:instr")) for field in fields] == ["PAGE"]


def test_render_book_docx_text_is_marked_italian() -> None:
    document = _open(render_book_docx(title="T", paragraphs=["x"]))

    normal_xml = document.styles["Normal"].element.xml
    assert 'w:val="it-IT"' in normal_xml


def test_render_book_docx_title_uses_body_serif_not_word_default() -> None:
    document = _open(render_book_docx(title="T", paragraphs=["x"]))

    assert document.styles["Title"].font.name == "Cambria"
    assert document.styles["Normal"].font.name == "Cambria"


def test_render_book_docx_without_paragraphs_keeps_only_the_title() -> None:
    document = _open(render_book_docx(title="Vuota", paragraphs=[]))

    assert [p.text for p in document.paragraphs] == ["Vuota"]
