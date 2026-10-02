from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import docx
import pytest
from pptx import Presentation
from pptx.util import Inches

from sbobina import document_extract
from sbobina.document_models import DocumentKind
from sbobina.document_sniff import ArchiveTooLargeError, check_archive_limits


def _write_presentation(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    ).text = "Safe text ENTITY_MARKER"
    presentation.save(str(path))


@pytest.mark.parametrize("external", [False, True])
def test_extract_pptx_entities_are_never_expanded(
    tmp_path: Path, external: bool
) -> None:
    path = tmp_path / "entity.pptx"
    _write_presentation(path=path)
    safe = document_extract.extract(path=path, kind=DocumentKind.PPTX)
    assert safe.pages[0].text == "Safe text ENTITY_MARKER"
    secret = tmp_path / "secret.txt"
    secret.write_text("PRIVATE_ENTITY_CONTENT", encoding="utf-8")
    value = f'SYSTEM "{secret.as_uri()}"' if external else '"PRIVATE_ENTITY_CONTENT"'
    with ZipFile(path) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    xml = entries["ppt/slides/slide1.xml"]
    header, body = xml.split(b"?>", maxsplit=1)
    entries["ppt/slides/slide1.xml"] = (
        header
        + b"?>"
        + f"<!DOCTYPE p:sld [<!ENTITY x {value}>]>".encode()
        + body.replace(b"ENTITY_MARKER", b"&x;")
    )
    with ZipFile(path, mode="w") as target:
        for name, data in entries.items():
            target.writestr(zinfo_or_arcname=name, data=data)
    result = document_extract.extract(path=path, kind=DocumentKind.PPTX)
    assert "Safe text" in result.pages[0].text
    assert "PRIVATE_ENTITY_CONTENT" not in result.pages[0].text


@pytest.mark.parametrize("kind", [DocumentKind.DOCX, DocumentKind.PPTX])
def test_extract_office_checks_archive_before_parser(
    tmp_path: Path, kind: DocumentKind
) -> None:
    path = tmp_path / f"document.{kind.value}"
    if kind == DocumentKind.DOCX:
        docx.Document().save(str(path))
    else:
        _write_presentation(path=path)
    calls: list[str] = []

    def check(path: Path) -> None:
        calls.append("check")
        check_archive_limits(path=path)

    parser_name = "Document" if kind == DocumentKind.DOCX else "Presentation"
    parser = getattr(document_extract, parser_name)

    def open_parser(path: Path) -> object:
        calls.append("parser")
        return parser(path)

    with (
        patch.object(document_extract, "check_archive_limits", check),
        patch.object(document_extract, parser_name, open_parser),
    ):
        document_extract.extract(path=path, kind=kind)
    assert calls == ["check", "parser"]
    with (
        patch.object(
            document_extract, "check_archive_limits", side_effect=ArchiveTooLargeError()
        ),
        patch.object(document_extract, parser_name) as blocked,
    ):
        with pytest.raises(ArchiveTooLargeError):
            document_extract.extract(path=path, kind=kind)
        blocked.assert_not_called()
