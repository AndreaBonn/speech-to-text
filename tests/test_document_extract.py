from pathlib import Path

import docx
import pytest
from document_fixtures import write_pdf
from pptx import Presentation
from pptx.util import Inches

from sbobina.document_extract import extract
from sbobina.document_models import DocumentKind, DocumentStatus
from sbobina.extracted_text import normalize_pages


def test_extract_pdf_three_pages_preserves_text(tmp_path: Path) -> None:
    path = tmp_path / "manual.pdf"
    texts = (
        "First page of the manual",
        "Second page about physics",
        "Third page with exercises",
    )
    write_pdf(path=path, texts=texts)
    result = extract(path=path, kind=DocumentKind.PDF)
    assert tuple(page.text for page in result.pages) == texts
    assert result.status == DocumentStatus.READY


def test_extract_pdf_images_marks_ready_no_text(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    write_pdf(path=path, texts=("", "", ""), image_only=True)
    scanned = extract(path=path, kind=DocumentKind.PDF)
    write_pdf(path=path, texts=("A readable page with sufficient text",))
    readable = extract(path=path, kind=DocumentKind.PDF)
    assert scanned.status == DocumentStatus.READY_NO_TEXT
    assert all(page.no_text for page in scanned.pages)
    assert readable.status == DocumentStatus.READY
    assert readable.pages[0].no_text is False


def test_extract_pptx_orders_shapes_and_includes_notes(tmp_path: Path) -> None:
    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    for title in ("First topic", "Second topic"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(3), Inches(5), Inches(1)
        ).text = "Body"
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(5), Inches(1)
        ).text = title
        notes = slide.notes_slide.notes_text_frame
        assert notes is not None
        notes.text = f"Notes for {title}"
    presentation.save(str(path))
    result = extract(path=path, kind=DocumentKind.PPTX)
    assert tuple(page.text for page in result.pages) == (
        "First topic\nBody\nNotes for First topic",
        "Second topic\nBody\nNotes for Second topic",
    )


@pytest.mark.parametrize("kind", [DocumentKind.TXT, DocumentKind.MD, DocumentKind.DOCX])
def test_extract_logical_pages_respects_headings_and_word_limit(
    tmp_path: Path, kind: DocumentKind
) -> None:
    path = tmp_path / f"notes.{kind.value}"
    if kind == DocumentKind.DOCX:
        document = docx.Document()
        document.add_heading("First", level=1)
        document.add_paragraph("word " * 600)
        document.add_heading("Second", level=1)
        document.add_paragraph("Conclusion")
        document.save(str(path))
    else:
        path.write_text(
            "# First\n" + "word " * 600 + "\n# Second\nConclusion", encoding="utf-8"
        )
    pages = extract(path=path, kind=kind).pages
    assert len(pages) == 3
    assert len(pages[0].text.split()) == 500
    assert pages[-1].text.endswith("Second\nConclusion")
    assert sum(page.text.split().count("word") for page in pages) == 600


def test_extract_cp1252_declares_fallback_and_normalizes(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_bytes("È una spiega-\nzione".encode("cp1252"))
    result = extract(path=path, kind=DocumentKind.TXT)
    assert result.encoding == "cp1252"
    assert result.pages[0].text == "È una spiegazione"


def test_normalize_pages_removes_repeated_edges_preserves_body() -> None:
    pages = normalize_pages(
        texts=("Header\nﬁsica\nFooter", "Header\ncorpo\nFooter", "Unique\ncorpo\nEnd")
    )
    assert pages == ("fisica", "corpo", "Unique\ncorpo\nEnd")
    assert normalize_pages(texts=("Keep",)) == ("Keep",)
