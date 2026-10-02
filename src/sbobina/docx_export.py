from collections.abc import Mapping
from io import BytesIO

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.styles import CT_Style
from docx.section import Section
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from sbobina.generation_models import (
    GenerationCitation,
    GenerationQuestion,
    GenerationRecord,
    SummarySection,
    SummarySentence,
)
from sbobina.generation_render import format_citation_source

PAGE_WIDTH = Cm(21.0)
PAGE_HEIGHT = Cm(29.7)
PAGE_MARGIN = Cm(2.5)
# Cambria is metric-compatible with LibreOffice's Caladea, so the layout
# holds when the file is opened outside Word.
BODY_FONT = "Cambria"
BODY_SIZE = Pt(12)
LINE_SPACING = 1.15
FIRST_LINE_INDENT = Cm(0.6)
LANGUAGE = "it-IT"
TITLE_SIZE = Pt(24)
# Word's default Title is blue Calibri Light, foreign to a serif book page.
TITLE_COLOR = RGBColor(0x22, 0x22, 0x22)
# a-d: GenerationQuestion enforces exactly 4 options for multiple_choice.
OPTION_LETTERS = "abcd"
DEFAULT_EXAM_TITLE = "Compito"
DEFAULT_SOLUTIONS_TITLE = "Soluzioni"
DEFAULT_SUMMARY_TITLE = "Riassunto"


def _set_language(style_element: CT_Style) -> None:
    run_properties = style_element.get_or_add_rPr()
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), LANGUAGE)
    run_properties.append(language)


def _add_page_number(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def _set_up_page(section: Section) -> None:
    section.page_width, section.page_height = PAGE_WIDTH, PAGE_HEIGHT
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, PAGE_MARGIN)


def _set_up_styles(document: DocxDocument) -> None:
    body = document.styles["Normal"]
    body.font.name = BODY_FONT
    body.font.size = BODY_SIZE
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.first_line_indent = FIRST_LINE_INDENT
    body.paragraph_format.line_spacing = LINE_SPACING
    body.paragraph_format.space_after = Pt(0)
    _set_language(style_element=body.element)
    # Justified Italian text without hyphenation leaves wide word gaps.
    document.settings.element.append(OxmlElement("w:autoHyphenation"))

    title_style = document.styles["Title"]
    title_style.font.name = BODY_FONT
    title_style.font.size = TITLE_SIZE
    title_style.font.color.rgb = TITLE_COLOR


def _new_titled_document(title: str) -> DocxDocument:
    """A4 document with the shared page/style setup and a Title heading."""
    document = Document()
    _set_up_page(section=document.sections[0])
    _set_up_styles(document=document)
    document.core_properties.title = title
    heading = document.add_paragraph(title, style="Title")
    heading.paragraph_format.first_line_indent = Cm(0)
    return document


def _save(document: DocxDocument) -> bytes:
    _add_page_number(paragraph=document.sections[0].footer.paragraphs[0])
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_book_docx(title: str, paragraphs: list[str]) -> bytes:
    """Lay out the transcript as a printable A4 book-style document.

    Parameters
    ----------
    title : str
        Document title, shown as the first heading and in the file metadata.
    paragraphs : list[str]
        Body paragraphs, already free of timestamps and review marks.

    Returns
    -------
    bytes
        The ``.docx`` file content.
    """
    document = _new_titled_document(title=title)
    for text in paragraphs:
        document.add_paragraph(text)
    return _save(document=document)


def _add_numbered_question(document: DocxDocument, index: int, text: str) -> None:
    paragraph = document.add_paragraph(f"{index}. {text}")
    paragraph.paragraph_format.first_line_indent = Cm(0)


def _add_lettered_options(document: DocxDocument, options: tuple[str, ...]) -> None:
    for letter, option in zip(OPTION_LETTERS, options, strict=True):
        paragraph = document.add_paragraph(f"{letter}) {option}")
        paragraph.paragraph_format.first_line_indent = Cm(0)


def render_exam_docx(generation: GenerationRecord) -> bytes:
    """Compito: numbered questions and lettered options, never solutions.

    Multiple_choice options are listed a-d with no marker on the correct
    one; open/oral questions have no options. Solutions and citations
    belong only to ``render_solutions_docx``.
    """
    document = _new_titled_document(title=generation.topic or DEFAULT_EXAM_TITLE)
    for index, question in enumerate(generation.questions, start=1):
        _add_numbered_question(document=document, index=index, text=question.question)
        if question.options:
            _add_lettered_options(document=document, options=question.options)
    return _save(document=document)


def _format_citation(
    citation: GenerationCitation, doc_filenames: Mapping[str, str]
) -> str:
    """Format as '«quote» (file, pagina N)' or '«quote» (lezione, mm:ss)'."""
    source = format_citation_source(citation=citation, doc_filenames=doc_filenames)
    return f"«{citation.quote}» ({source})"


def _add_solution(
    document: DocxDocument,
    index: int,
    question: GenerationQuestion,
    doc_filenames: Mapping[str, str],
) -> None:
    _add_numbered_question(document=document, index=index, text=question.question)
    if question.options and question.correct_index is not None:
        letter = OPTION_LETTERS[question.correct_index]
        document.add_paragraph(f"Risposta corretta: {letter}")
    document.add_paragraph(f"Soluzione: {question.solution}")
    for citation in question.citations:
        document.add_paragraph(
            _format_citation(citation=citation, doc_filenames=doc_filenames)
        )


def render_solutions_docx(
    generation: GenerationRecord, doc_filenames: Mapping[str, str] | None = None
) -> bytes:
    """Soluzioni: per domanda la soluzione, la lettera corretta se a
    crocette, e le citazioni come «testo» (riferimento). Sempre in un file
    separato dal compito."""
    title = (
        f"{DEFAULT_SOLUTIONS_TITLE} - {generation.topic}"
        if generation.topic
        else DEFAULT_SOLUTIONS_TITLE
    )
    document = _new_titled_document(title=title)
    for index, question in enumerate(generation.questions, start=1):
        _add_solution(
            document=document,
            index=index,
            question=question,
            doc_filenames=doc_filenames or {},
        )
    return _save(document=document)


def _format_sentence(
    sentence: SummarySentence, doc_filenames: Mapping[str, str]
) -> str:
    if not sentence.citations:
        return sentence.text
    refs = "; ".join(
        format_citation_source(citation=citation, doc_filenames=doc_filenames)
        for citation in sentence.citations
    )
    return f"{sentence.text} ({refs})"


def _add_summary_section(
    document: DocxDocument, section: SummarySection, doc_filenames: Mapping[str, str]
) -> None:
    document.add_heading(section.title, level=1)
    for sentence in section.sentences:
        document.add_paragraph(
            _format_sentence(sentence=sentence, doc_filenames=doc_filenames)
        )


def render_summary_docx(
    generation: GenerationRecord, doc_filenames: Mapping[str, str] | None = None
) -> bytes:
    """Riassunto: sezioni come titoli, frasi come paragrafi, citazioni fra
    parentesi accanto alla frase che le porta."""
    document = _new_titled_document(title=generation.topic or DEFAULT_SUMMARY_TITLE)
    for section in generation.sections:
        _add_summary_section(
            document=document, section=section, doc_filenames=doc_filenames or {}
        )
    return _save(document=document)
