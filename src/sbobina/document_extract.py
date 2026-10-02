"""Single boundary for the three office/text parsers plus PDF (D4).

``extract`` is the only place that opens python-docx/python-pptx/pypdfium2
documents; callers never touch those libraries directly.
"""

from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.slide import Slide

from sbobina.document_models import DocumentKind, DocumentStatus
from sbobina.document_sniff import check_archive_limits
from sbobina.extracted_text import ExtractedText, Page, normalize_pages
from sbobina.pdf_text import extract_pdf_pages

LOGICAL_PAGE_WORD_LIMIT = 500
PAGE_NO_TEXT_CHAR_THRESHOLD = 20
NO_TEXT_PAGE_MAJORITY = 0.5
TEXT_FALLBACK_ENCODING = "cp1252"


def extract(path: Path, kind: DocumentKind) -> ExtractedText:
    """Extract page/slide text for a document already identified as ``kind``."""
    raw_pages, encoding = _raw_pages(path=path, kind=kind)
    pages = tuple(_build_page(text=text) for text in normalize_pages(texts=raw_pages))
    return ExtractedText(
        pages=pages, status=_status_from_pages(pages=pages), encoding=encoding
    )


def _build_page(text: str) -> Page:
    return Page(text=text, no_text=len(text) < PAGE_NO_TEXT_CHAR_THRESHOLD)


def _status_from_pages(pages: tuple[Page, ...]) -> DocumentStatus:
    if not pages:
        return DocumentStatus.READY_NO_TEXT
    no_text_ratio = sum(page.no_text for page in pages) / len(pages)
    if no_text_ratio > NO_TEXT_PAGE_MAJORITY:
        return DocumentStatus.READY_NO_TEXT
    return DocumentStatus.READY


def _raw_pages(path: Path, kind: DocumentKind) -> tuple[tuple[str, ...], str | None]:
    match kind:
        case DocumentKind.PDF:
            return extract_pdf_pages(path=path), None
        case DocumentKind.PPTX:
            return _extract_pptx(path=path), None
        case DocumentKind.DOCX:
            return _extract_docx(path=path), None
        case DocumentKind.TXT | DocumentKind.MD:
            return _extract_plain_text(path=path)


def _extract_pptx(path: Path) -> tuple[str, ...]:
    check_archive_limits(path=path)
    presentation = Presentation(str(path))
    return tuple(_slide_text(slide=slide) for slide in presentation.slides)


def _text_shapes(slide: Slide) -> list[Shape]:
    return [
        shape
        for shape in slide.shapes
        if isinstance(shape, Shape) and shape.has_text_frame and shape.text_frame.text
    ]


def _slide_text(slide: Slide) -> str:
    shapes = sorted(_text_shapes(slide=slide), key=lambda shape: shape.top or 0)
    lines = [shape.text_frame.text for shape in shapes]
    notes_frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
    if notes_frame is not None and notes_frame.text:
        lines.append(notes_frame.text)
    return "\n".join(lines)


def _extract_docx(path: Path) -> tuple[str, ...]:
    check_archive_limits(path=path)
    document = Document(str(path))
    items = [
        (
            bool(paragraph.style and paragraph.style.name.startswith("Heading")),
            paragraph.text,
        )
        for paragraph in document.paragraphs
    ]
    return _sections_to_pages(sections=_split_sections(items=items))


def _extract_plain_text(path: Path) -> tuple[tuple[str, ...], str]:
    raw = path.read_bytes()
    try:
        text, encoding = raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        text, encoding = raw.decode(TEXT_FALLBACK_ENCODING), TEXT_FALLBACK_ENCODING
    items = [
        (line.startswith("#"), line.lstrip("#").strip()) for line in text.split("\n")
    ]
    return _sections_to_pages(sections=_split_sections(items=items)), encoding


def _split_sections(
    items: list[tuple[bool, str]],
) -> list[tuple[str | None, list[str]]]:
    sections: list[tuple[str | None, list[str]]] = []
    heading: str | None = None
    body: list[str] = []
    for is_heading, text in items:
        if is_heading:
            if heading is not None or body:
                sections.append((heading, body))
            heading, body = text, []
        else:
            body.append(text)
    sections.append((heading, body))
    return sections


def _sections_to_pages(sections: list[tuple[str | None, list[str]]]) -> tuple[str, ...]:
    pages: list[str] = []
    for heading, body in sections:
        pages.extend(_section_to_page_texts(heading=heading, body=body))
    return tuple(pages)


def _section_to_page_texts(heading: str | None, body: list[str]) -> list[str]:
    lines = [heading, *body] if heading is not None else body
    natural = "\n".join(lines)
    words = natural.split()
    if len(words) <= LOGICAL_PAGE_WORD_LIMIT:
        return [natural]
    return [
        " ".join(words[start : start + LOGICAL_PAGE_WORD_LIMIT])
        for start in range(0, len(words), LOGICAL_PAGE_WORD_LIMIT)
    ]
